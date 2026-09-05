#!/usr/bin/env python3
import openpyxl
import os
import uuid
from datetime import datetime

import json

EXCEL_PATH = '/home/dns/Downloads/MALL-APP-MIGRACAO-REACT-NATIVE/Mapa - Sinalização do Mall.xlsx'
SYNC_EXPORT_DIR = '/home/dns/Desenvolvimento/Sistema_Centro_Fashion/database/sync_export'
OUTPUT_SQL = '/home/dns/Desenvolvimento/Sistema_Centro_Fashion/database/002_seed_data.sql'

def escape_sql(val):
    if val is None:
        return 'NULL'
    val_str = str(val).replace("'", "''")
    return f"'{val_str}'"

def get_table_rows(sheet_name, wb=None):
    json_path = os.path.join(SYNC_EXPORT_DIR, f"{sheet_name}.json")
    if os.path.exists(json_path):
        with open(json_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    if wb and sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        return list(sheet.iter_rows(values_only=True))
    return []

def generate_seed():
    wb = None
    if os.path.exists(EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
        except Exception:
            wb = None

    sql_lines = [
        "-- ==============================================================================",
        "-- DADOS DE SEED E MIGRAÇÃO DO LEGADO (MAPA - SINALIZAÇÃO DO MALL)",
        f"-- Gerado em: {datetime.utcnow().isoformat()}Z",
        "-- ==============================================================================\n",
    ]

    # 1. MIGRAÇÃO DE SETORES
    sector_rows = get_table_rows('SETORES', wb)
    if sector_rows:
        sql_lines.append("-- 1. SETORES")
        for row in sector_rows[1:]:
            if not row or not row[0]:
                continue
            sector_code = row[3] if len(row) > 3 and row[3] else row[0]
            name = row[2] if len(row) > 2 and row[2] else sector_code
            color = row[4] if len(row) > 4 and row[4] else '#000000'
            sql = f"""INSERT INTO sectors (code, name, color_hex)
VALUES ({escape_sql(sector_code)}, {escape_sql(name)}, {escape_sql(color)})
ON CONFLICT (code) DO UPDATE SET name = EXCLUDED.name, color_hex = EXCLUDED.color_hex;"""
            sql_lines.append(sql)
        sql_lines.append("")

    # 2. MIGRAÇÃO DE USUÁRIOS
    user_rows = get_table_rows('USUARIOS', wb)
    if user_rows:
        sql_lines.append("-- 2. USUÁRIOS")
        for row in user_rows[1:]:
            if not row or not row[0]:
                continue
            email = row[0]
            name = row[1] or email
            role = str(row[2] or 'INSPECTOR').upper()
            active = 'true' if row[3] is True or str(row[3]).upper() == 'SIM' or str(row[3]).upper() == 'TRUE' else 'false'
            sql = f"""INSERT INTO users (email, name, profile_role, active)
VALUES ({escape_sql(email)}, {escape_sql(name)}, {escape_sql(role)}, {active})
ON CONFLICT (email) DO UPDATE SET name = EXCLUDED.name, profile_role = EXCLUDED.profile_role;"""
            sql_lines.append(sql)
        sql_lines.append("")

    # 3. MIGRAÇÃO DE SINALIZAÇÕES (REGISTROS)
    signage_map = {} # legacy_id -> uuid
    registro_rows = get_table_rows('REGISTROS', wb)
    if registro_rows:
        sql_lines.append("-- 3. SINALIZAÇÕES DE CAMPO (REGISTROS)")
        for row in registro_rows[1:]:
            if not row or not row[0]:
                continue
            legacy_id = str(row[0])
            protocol = row[2] or legacy_id
            status = row[5] or 'ATIVA'
            lifecycle_status = 'ACTIVE' if status == 'ATIVA' else 'DECOMMISSIONED'
            category = row[6] or 'Placa'
            notes = str(row[9]) if len(row) > 9 and row[9] else str(row[8]) if len(row) > 8 and row[8] else ''
            
            asset_uuid = str(uuid.uuid4())
            signage_map[legacy_id] = asset_uuid

            sql = f"""INSERT INTO signage_assets (id, legacy_id, asset_code, category, conservation_status, lifecycle_status, notes)
VALUES ('{asset_uuid}', {escape_sql(legacy_id)}, {escape_sql(protocol)}, {escape_sql(category)}, 'GOOD', {escape_sql(lifecycle_status)}, {escape_sql(notes)})
ON CONFLICT (asset_code) DO NOTHING;"""
            sql_lines.append(sql)
        sql_lines.append("")

    # 4. MIGRAÇÃO DE FOTOS DE REGISTRO
    foto_rows = get_table_rows('REGISTRO_FOTOS', wb)
    if foto_rows:
        sql_lines.append("-- 4. FOTOS E MÍDIAS DE CAMPO")
        for row in foto_rows[1:]:
            if not row or not row[0]:
                continue
            legacy_photo_id = str(row[0])
            legacy_registro_id = str(row[3]) if len(row) > 3 else None
            drive_file_id = str(row[7]) if len(row) > 7 and row[7] else None
            filename = str(row[8]) if len(row) > 8 and row[8] else f"{legacy_photo_id}.jpg"
            mime_type = str(row[9]) if len(row) > 9 and row[9] else 'image/jpeg'
            
            media_uuid = str(uuid.uuid4())
            storage_key = f"photos/migrated/{filename}"

            sql_media = f"""INSERT INTO media_assets (id, legacy_drive_id, storage_key, mime_type, size_bytes, sha256)
VALUES ('{media_uuid}', {escape_sql(drive_file_id)}, {escape_sql(storage_key)}, {escape_sql(mime_type)}, 1024, {escape_sql(media_uuid.replace('-', ''))})
ON CONFLICT DO NOTHING;"""
            sql_lines.append(sql_media)

            if legacy_registro_id and legacy_registro_id in signage_map:
                signage_uuid = signage_map[legacy_registro_id]
                sql_link = f"""INSERT INTO signage_media (signage_id, media_id, is_primary)
VALUES ('{signage_uuid}', '{media_uuid}', true)
ON CONFLICT DO NOTHING;"""
                sql_lines.append(sql_link)

        sql_lines.append("")

    with open(OUTPUT_SQL, 'w', encoding='utf-8') as f:
        f.write('\n'.join(sql_lines))

    print(f"Seed SQL gerado com sucesso em: {OUTPUT_SQL}")

if __name__ == '__main__':
    generate_seed()
